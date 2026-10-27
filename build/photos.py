#!/usr/bin/env python3
"""Bring new photos into images/GFC-<code>/ as JPEG (long side max 2048 px):
 1) files Mariia uploaded on GitHub into photos_in/GFC-<code>/ (any jpg/jpeg/png/webp) -> next free NN.jpg, upload removed;
 2) rows of tab «Фото» with «Показувати» = так, a file that does not exist yet and a link in «Посилання для завантаження»
    (direct image URL or a Google Drive file link shared as «Anyone with the link») -> downloaded under that «Файл» name.
Photos already in images/ are never changed."""
import argparse, io, os, re, sys, urllib.request
import openpyxl
from PIL import Image, ImageOps
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sheet_config import refid, YES

MAX_SIDE, Q = 2048, 88
EXT = (".jpg", ".jpeg", ".png", ".webp")

def save_jpeg(data, dst):
    im = Image.open(io.BytesIO(data)); im = ImageOps.exif_transpose(im).convert("RGB")
    if max(im.size) > MAX_SIDE: im.thumbnail((MAX_SIDE, MAX_SIDE))
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    im.save(dst, "JPEG", quality=Q, optimize=True)
    return im.size

def next_name(d):
    nums = [int(m.group(1)) for f in (os.listdir(d) if os.path.isdir(d) else []) if (m := re.fullmatch(r"(\d+)\.jpg", f))]
    return f"{(max(nums) + 1) if nums else 1:02d}.jpg"

def drive(url):
    m = re.search(r"/file/d/([\w-]+)", url) or re.search(r"[?&]id=([\w-]+)", url)
    return f"https://drive.google.com/uc?export=download&id={m.group(1)}" if m and "drive.google.com" in url else url

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--xlsx", required=True); ap.add_argument("--root", default=".")
    a = ap.parse_args(); img = os.path.join(a.root, "images"); inbox = os.path.join(a.root, "photos_in"); log = []
    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    known = {refid(r[0]) for r in wb["Параметри"].iter_rows(min_row=2, values_only=True)} - {None} if "Параметри" in wb.sheetnames else set()
    # 1) photos_in/
    if os.path.isdir(inbox):
        for d in sorted(os.listdir(inbox)):
            src = os.path.join(inbox, d)
            if not os.path.isdir(src): continue
            k = refid(d)
            if not k or k not in known:
                log.append(f"photos_in/{d}: папка має називатися як Realting ID з вкладки «Параметри» (напр. GFC-CAPRI або GFC-3843771) – пропущено"); continue
            for f in sorted(os.listdir(src)):
                p = os.path.join(src, f)
                if not f.lower().endswith(EXT): continue
                dst = os.path.join(img, f"GFC-{k}", next_name(os.path.join(img, f"GFC-{k}")))
                try: w, h = save_jpeg(open(p, "rb").read(), dst)
                except Exception as e: log.append(f"photos_in/{d}/{f}: не вдалося відкрити як фото ({type(e).__name__})"); continue
                os.remove(p); log.append(f"GFC-{k}: {f} -> {os.path.basename(dst)} ({w}×{h})")
    # 2) links in tab «Фото»
    if "Фото" in wb.sheetnames:
        ws = wb["Фото"]; hdr = [str(c.value).strip() if c.value else "" for c in ws[1]]
        col = {h: i for i, h in enumerate(hdr)}
        L = col.get("Посилання для завантаження")
        for n, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            k = refid(r[col["Realting ID"]]); url = (r[L] if L is not None and L < len(r) else None)
            if not k or not url or str(r[col["Показувати"]] or "").strip().lower() not in YES: continue
            fn = str(r[col["Файл"]] or "").strip()
            d = os.path.join(img, f"GFC-{k}")
            if fn and os.path.exists(os.path.join(d, fn)): continue
            fn = fn if re.fullmatch(r"[\w-]+\.jpg", fn or "") else next_name(d)
            try:
                data = urllib.request.urlopen(urllib.request.Request(drive(str(url).strip()), headers={"User-Agent": "Mozilla/5.0"}), timeout=60).read()
                w, h = save_jpeg(data, os.path.join(d, fn)); log.append(f"GFC-{k}: завантажено {fn} ({w}×{h}) з вкладки «Фото», рядок {n}")
            except Exception as e:
                log.append(f"GFC-{k}: «Фото» рядок {n} – не вдалося завантажити фото ({type(e).__name__}). Посилання має відкривати саме фото, Drive – доступ «Усі, хто має посилання».")
    for l in log: print("PHOTO", l)

if __name__ == "__main__":
    main()
