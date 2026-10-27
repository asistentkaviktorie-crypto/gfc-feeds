"""Read the per-listing settings that used to live in local files from the sheet tabs:
  «Параметри» -> META (town, coordinates, Prian codes, types, features …)   [was build_feed.py META]
  «Сайт»      -> SITE list (site-only fields, order = row order)            [was site_feed/meta.py]
  «Переклади» -> TXT[lang][key] = title/summary/lifestyle/invest/risk       [was site_feed/<lang>.py]
  «Фото»      -> PHOTOS[key] = ordered photos, sizes read from images/GFC-<key>/<file> [was images/manifest.json]
Any problem that would publish wrong data raises SystemExit with a Ukrainian message (the workflow then stops)."""
import os, re
import openpyxl

YES = {"так", "yes", "y", "1", "true", "да", "+"}
LANGS = ["en", "es", "ru", "de", "fr", "cs"]
TFIELDS = [("Назва", "title"), ("Коротко", "summary"), ("Стиль життя", "lifestyle"), ("Інвестиції", "invest"), ("Ризики", "risk")]

def refid(v):
    if v in (None, ""): return None
    s = str(v).strip()
    try:
        f = float(s.replace(",", "."))
        if f.is_integer() and f > 0: return str(int(f))
    except ValueError: pass
    s = s.upper()
    if s.startswith("GFC-"): s = s[4:]
    return s if re.fullmatch(r"[A-Z0-9]+", s) else None

def blank(v): return v is None or (isinstance(v, str) and v.strip() == "")

def num(v, where):
    if blank(v): return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v) if float(v).is_integer() else float(v)
    s = str(v).strip().replace("\u00a0", "").replace(" ", "")
    pct = s.endswith("%"); s = s.rstrip("%").replace(",", ".")
    try: f = float(s)
    except ValueError: raise SystemExit(f"ПОМИЛКА {where}: «{v}» – тут має бути число")
    if pct: f /= 100
    return int(f) if f.is_integer() else f

def frac(v, where):  # 90% / 0.9 / 90 -> 0.9
    n = num(v, where)
    if n is None: return None
    n = n / 100 if n > 1 else n
    return round(n, 6)

def text(v): return None if blank(v) else str(v).strip()

def table(wb, name, required):
    if name not in wb.sheetnames: raise SystemExit(f"ПОМИЛКА: у таблиці немає вкладки «{name}»")
    ws = wb[name]
    hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    miss = [h for h in required if h not in hdr]
    if miss: raise SystemExit(f"ПОМИЛКА: у вкладці «{name}» немає колонок {miss} (не перейменовуйте заголовки)")
    out = []
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        d = {h: (r[j] if j < len(r) else None) for j, h in enumerate(hdr) if h}
        k = refid(d.get("Realting ID"))
        if k: d["_key"] = k; d["_row"] = i; out.append(d)
    return out

def load_meta(wb):
    req = ["Realting ID", "Місто", "Провінція", "Локація детально", "Prian район (код)", "Адреса EN", "Адреса RU", "Широта", "Довгота",
           "Тип Kyero", "Тип Prian (код)", "Новобудова (1/0)", "Prian стан (код)", "Рік побудови", "Басейн (1/0)", "Енергоклас",
           "Підпис EN (до 40 знаків)", "Підпис RU (до 40 знаків)", "Особливості EN (через ;)"]
    META = {}
    for d in table(wb, "Параметри", req):
        w = f"«Параметри» рядок {d['_row']} ({d['Realting ID']})"
        for h in ["Місто", "Провінція", "Prian район (код)", "Широта", "Довгота", "Тип Kyero", "Тип Prian (код)", "Підпис EN (до 40 знаків)", "Підпис RU (до 40 знаків)"]:
            if blank(d[h]): raise SystemExit(f"ПОМИЛКА {w}: порожня колонка «{h}»")
        m = dict(town=text(d["Місто"]), province=text(d["Провінція"]), detail=text(d["Локація детально"]) or "",
                 district_id=int(num(d["Prian район (код)"], w)), addr_en=text(d["Адреса EN"]) or "", addr_ru=text(d["Адреса RU"]) or "",
                 lat=float(num(d["Широта"], w)), lon=float(num(d["Довгота"], w)), kt=text(d["Тип Kyero"]).lower(),
                 pt=int(num(d["Тип Prian (код)"], w)), nb=int(num(d["Новобудова (1/0)"], w) or 0),
                 bt=None if blank(d["Prian стан (код)"]) else int(num(d["Prian стан (код)"], w)),
                 pool=int(num(d["Басейн (1/0)"], w) or 0), st_en=text(d["Підпис EN (до 40 знаків)"]), st_ru=text(d["Підпис RU (до 40 знаків)"]))
        if not blank(d["Рік побудови"]): m["bdate"] = str(num(d["Рік побудови"], w))
        if not blank(d["Енергоклас"]): m["energy"] = text(d["Енергоклас"]).upper()
        if not blank(d["Особливості EN (через ;)"]):
            m["features"] = [f.strip() for f in str(d["Особливості EN (через ;)"]).split(";") if f.strip()]
        if not (-90 <= m["lat"] <= 90 and -180 <= m["lon"] <= 180): raise SystemExit(f"ПОМИЛКА {w}: неправильні координати")
        if d["_key"] in META: raise SystemExit(f"ПОМИЛКА {w}: цей Realting ID уже є вище у «Параметри»")
        META[d["_key"]] = m
    return META

def load_site(wb):
    req = ["Realting ID", "ID на сайті", "ID кабінету", "Тип", "Район (area)", "Місто", "Ліцензія", "Комуналка €/міс", "IBI €/рік",
           "Страховка €/рік", "До моря, м", "Довга оренда €/міс", "Ціна за ніч (ADR) €", "Заповнюваність довга", "Заповнюваність коротка",
           "Управління довга", "Управління коротка", "Ціна «від»", "Приховати на сайті", "Обкладинка (файл)"]
    SITE, COVER = [], {}
    for d in table(wb, "Сайт", req):
        w = f"«Сайт» рядок {d['_row']} ({d['Realting ID']})"
        k = d["_key"]
        typ = (text(d["Тип"]) or "").lower()
        if typ not in ("new", "resale", "land"): raise SystemExit(f"ПОМИЛКА {w}: «Тип» має бути new / resale / land")
        SITE.append(dict(key=k, id=text(d["ID на сайті"]) or f"gfc-{k.lower()}", cabinetId=text(d["ID кабінету"]), type=typ,
            area=text(d["Район (area)"]), town=text(d["Місто"]), license=(text(d["Ліцензія"]) or "no").lower(),
            communityMonth=num(d["Комуналка €/міс"], w), ibiYear=num(d["IBI €/рік"], w), insuranceYear=num(d["Страховка €/рік"], w),
            seaM=num(d["До моря, м"], w), longRentMonth=num(d["Довга оренда €/міс"], w), adr=num(d["Ціна за ніч (ADR) €"], w),
            occLong=frac(d["Заповнюваність довга"], w), occShort=frac(d["Заповнюваність коротка"], w),
            mgmtLong=frac(d["Управління довга"], w), mgmtShort=frac(d["Управління коротка"], w),
            priceFrom=str(d["Ціна «від»"] or "").strip().lower() in YES, hidden=str(d["Приховати на сайті"] or "").strip().lower() in YES))
        if not blank(d["Обкладинка (файл)"]): COVER[k] = text(d["Обкладинка (файл)"])
    return SITE, COVER

def load_texts(wb):
    TXT = {L: {} for L in LANGS}
    for d in table(wb, "Переклади", ["Realting ID", "Мова"] + [h for h, _ in TFIELDS]):
        L = (text(d["Мова"]) or "").lower()
        if L not in TXT: raise SystemExit(f"ПОМИЛКА «Переклади» рядок {d['_row']}: мова «{d['Мова']}» – дозволено {LANGS}")
        TXT[L][d["_key"]] = {f: (text(d[h]) or "") for h, f in TFIELDS}
    return TXT

def load_photos(wb, img_dir, warn):
    from PIL import Image
    rows = table(wb, "Фото", ["Realting ID", "№", "Файл", "Показувати"])
    by = {}
    for d in rows: by.setdefault(d["_key"], []).append(d)
    PHOTOS = {}
    keys = list(by) + [k[4:] for k in sorted(os.listdir(img_dir)) if k.startswith("GFC-") and k[4:] not in by] if os.path.isdir(img_dir) else list(by)
    for k in keys:
        lst, listed = [], set()
        rs = by.get(k, [])
        rs = sorted(rs, key=lambda d: (num(d["№"], "Фото") if not blank(d["№"]) else 10**6, d["_row"]))
        for d in rs:
            fn = text(d["Файл"]) or ""
            listed.add(fn)
            if str(d["Показувати"] or "").strip().lower() not in YES: continue
            p = os.path.join(img_dir, f"GFC-{k}", fn)
            if not fn or not os.path.exists(p):
                warn.append(f"GFC-{k}: фото «{fn}» (вкладка «Фото», рядок {d['_row']}) немає в images/GFC-{k}/ – пропущено"); continue
            with Image.open(p) as im: wdt, hgt = im.size
            lst.append(dict(n=os.path.splitext(fn)[0], w=wdt, h=hgt))
        d_ = os.path.join(img_dir, f"GFC-{k}")
        if os.path.isdir(d_):
            for fn in sorted(os.listdir(d_)):
                if fn.lower().endswith(".jpg") and fn not in listed:
                    with Image.open(os.path.join(d_, fn)) as im: wdt, hgt = im.size
                    lst.append(dict(n=os.path.splitext(fn)[0], w=wdt, h=hgt))
                    if rs: warn.append(f"GFC-{k}: нове фото {fn} додано в кінець – впишіть його у вкладку «Фото», щоб задати порядок")
        PHOTOS[k] = lst
    return PHOTOS

def load(xlsx, img_dir, warn):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    META = load_meta(wb); SITE, COVER = load_site(wb); TXT = load_texts(wb)
    PHOTOS = load_photos(wb, img_dir, warn)
    return dict(META=META, SITE=SITE, COVER=COVER, TXT=TXT, PHOTOS=PHOTOS)
